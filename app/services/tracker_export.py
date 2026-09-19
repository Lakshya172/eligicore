"""Tracker export — client-supplied tracking rows rendered as an ``.xlsx`` workbook (ADR-023).

A renderer, not a tracking service: ``TrackerRecord[]`` in, workbook bytes out. It reads no
database, calls no eligibility, matching, recommendation or AI code, and never corrects or
enriches a row. What the client sends is what the workbook says.

**Nothing leaves memory.** The workbook is built and saved into a ``BytesIO`` buffer — no
temporary file, no cache — and the only log line carries counts (INV-1, INV-4).

**No cell can execute.** Every string goes through :func:`_write_text`, which neutralises a
leading formula character and pins the cell type to text (ADR-023 §6).

Framework-free: importable and testable without FastAPI (INV-7).
"""

from __future__ import annotations

import logging
import time
from collections.abc import Callable, Sequence
from datetime import date, datetime, timezone
from enum import Enum
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.eligibility import RequirementResult
from app.schemas.tracker import TrackerRecord

logger = logging.getLogger("eligicore.tracker_export")

#: Media type of the returned file.
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: Download filename. Fixed: no candidate data, no timestamp (ADR-023 §7).
TRACKER_FILENAME = "eligicore-tracker.xlsx"

TRACKER_SHEET = "Tracker"
REQUIREMENTS_SHEET = "Requirements"

#: Fixed workbook properties. Nothing here may come from the request (ADR-023 §7).
WORKBOOK_CREATOR = "EligiCore"
WORKBOOK_TITLE = "EligiCore application tracker"

#: A text value starting with one of these could be read as a formula by a spreadsheet.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

#: Prepended to a neutralised value: keeps it readable and marks it as literal text.
TEXT_MARKER = "'"

DATE_FORMAT = "yyyy-mm-dd"
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"

CellValue = str | int | float | date | datetime | Enum | None

#: ``Tracker`` sheet: header, column width, and how a record fills the column. The order is the
#: published layout — change it only by adding columns at the end (ADR-010 by analogy).
_TRACKER_COLUMNS: tuple[tuple[str, int, Callable[[int, TrackerRecord], CellValue]], ...] = (
    ("Row", 6, lambda row, record: row),
    ("Job ID", 20, lambda row, record: record.job_id),
    ("Company", 28, lambda row, record: record.company_name),
    ("Role", 32, lambda row, record: record.role_title),
    ("Job Status", 12, lambda row, record: record.job_status),
    ("Eligibility", 18, lambda row, record: record.eligibility_state),
    ("Match Score", 12, lambda row, record: record.match_score),
    ("Application Status", 18, lambda row, record: record.application_status),
    ("Deadline", 12, lambda row, record: record.deadline),
    ("Apply Link", 40, lambda row, record: record.apply_link),
    ("Reason", 60, lambda row, record: record.reason),
    ("Evaluated At (UTC)", 20, lambda row, record: record.evaluated_at),
    ("Notes", 40, lambda row, record: record.notes),
)

#: ``Requirements`` sheet. ``Tracker Row`` and ``Job ID`` tie each entry to its parent row.
_REQUIREMENT_COLUMNS: tuple[
    tuple[str, int, Callable[[int, TrackerRecord, RequirementResult], CellValue]], ...
] = (
    ("Tracker Row", 11, lambda row, record, req: row),
    ("Job ID", 20, lambda row, record, req: record.job_id),
    ("Requirement Type", 20, lambda row, record, req: req.requirement_type),
    ("Requirement", 40, lambda row, record, req: req.requirement),
    ("Candidate Value", 24, lambda row, record, req: req.candidate_value),
    ("Status", 10, lambda row, record, req: req.status),
    ("Confidence", 12, lambda row, record, req: req.confidence),
    ("Method", 14, lambda row, record, req: req.method),
    ("Reason Code", 28, lambda row, record, req: req.reason_code),
    ("Note", 60, lambda row, record, req: req.note),
)

TRACKER_HEADERS = tuple(header for header, _, _ in _TRACKER_COLUMNS)
REQUIREMENT_HEADERS = tuple(header for header, _, _ in _REQUIREMENT_COLUMNS)


# ---------------------------------------------------------------------------------------
# Cells
# ---------------------------------------------------------------------------------------


def neutralize(text: str) -> str:
    """Return ``text`` so that no spreadsheet reads it as a formula.

    A value starting with ``=``, ``+``, ``-``, ``@``, tab or carriage return gets a leading
    apostrophe; anything else is returned unchanged. Deterministic and visible: ``=1+1`` becomes
    ``'=1+1``.
    """
    return TEXT_MARKER + text if text.startswith(FORMULA_PREFIXES) else text


def _write_text(cell: Cell, text: str) -> None:
    """Write ``text`` as literal text — neutralised, and with the cell type pinned to string.

    Two independent guards: the prefix protects the value in any spreadsheet, and the explicit
    type stops openpyxl from storing a string that begins with ``=`` as a formula.
    """
    cell.value = neutralize(text)
    cell.data_type = "s"


def _write(cell: Cell, value: CellValue) -> None:
    """Write one value with the cell type its meaning calls for. ``None`` stays empty."""
    if value is None:
        return
    if isinstance(value, Enum):
        _write_text(cell, str(value.value))
    elif isinstance(value, str):
        _write_text(cell, value)
    elif isinstance(value, datetime):
        # Excel has no time zones: store the UTC wall-clock time, labelled in the header.
        cell.value = value.astimezone(timezone.utc).replace(tzinfo=None)
        cell.number_format = DATETIME_FORMAT
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = DATE_FORMAT
    else:
        cell.value = value


def _write_header(sheet: Worksheet, columns: Sequence[tuple[str, int, object]]) -> None:
    bold = Font(bold=True)
    for index, (header, width, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        _write_text(cell, header)
        cell.font = bold
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


# ---------------------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------------------


def build_workbook(records: Sequence[TrackerRecord]) -> Workbook:
    """Lay the records out as a workbook, in the order given."""
    workbook = Workbook()
    properties = workbook.properties
    properties.creator = WORKBOOK_CREATOR
    properties.lastModifiedBy = WORKBOOK_CREATOR
    properties.title = WORKBOOK_TITLE

    tracker = workbook.active
    tracker.title = TRACKER_SHEET
    _write_header(tracker, _TRACKER_COLUMNS)
    for row, record in enumerate(records, start=1):
        for column, (_, _, value_of) in enumerate(_TRACKER_COLUMNS, start=1):
            _write(tracker.cell(row=row + 1, column=column), value_of(row, record))

    requirement_rows = [
        (row, record, requirement)
        for row, record in enumerate(records, start=1)
        for requirement in record.requirement_breakdown or []
    ]
    if requirement_rows:
        sheet = workbook.create_sheet(REQUIREMENTS_SHEET)
        _write_header(sheet, _REQUIREMENT_COLUMNS)
        for line, (row, record, requirement) in enumerate(requirement_rows, start=2):
            for column, (_, _, value_of) in enumerate(_REQUIREMENT_COLUMNS, start=1):
                _write(sheet.cell(row=line, column=column), value_of(row, record, requirement))

    return workbook


def export_tracker(records: Sequence[TrackerRecord]) -> bytes:
    """Render tracker rows as ``.xlsx`` bytes, entirely in memory.

    Logs one line of counts. Never logs a job id, company, role, link, reason, note,
    requirement, score or any candidate value.
    """
    started = time.perf_counter()
    workbook = build_workbook(records)
    buffer = BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()

    logger.info(
        "tracker_export rows=%d requirement_rows=%d bytes=%d duration_ms=%.1f",
        len(records),
        sum(len(record.requirement_breakdown or []) for record in records),
        len(content),
        (time.perf_counter() - started) * 1000,
    )
    return content
