"""Tracker export — schema, workbook, formula-injection and architecture tests (PR 6A, ADR-023).

Every workbook is reopened with openpyxl and, where it matters, read as raw sheet XML: the file
is the contract, not the objects that built it. All data is obviously synthetic.
"""

from __future__ import annotations

import ast
import builtins
import io
import os
import pathlib
import tempfile
import zipfile
from datetime import date, datetime, timezone
from typing import Any

import openpyxl
import pytest
from pydantic import ValidationError

from app.schemas.tracker import (
    MAX_REQUIREMENT_TEXT,
    MAX_REQUIREMENTS_PER_ROW,
    MAX_TRACKER_ROWS,
    ApplicationStatus,
    TrackerExportRequest,
    TrackerRecord,
)
from app.services import tracker_export
from app.services.tracker_export import (
    REQUIREMENT_HEADERS,
    REQUIREMENTS_SHEET,
    TRACKER_FILENAME,
    TRACKER_HEADERS,
    TRACKER_SHEET,
    WORKBOOK_CREATOR,
    WORKBOOK_TITLE,
    build_workbook,
    export_tracker,
    neutralize,
)

SERVICE_PATH = pathlib.Path(tracker_export.__file__)

EXPECTED_TRACKER_HEADERS = (
    "Row", "Job ID", "Company", "Role", "Job Status", "Eligibility", "Match Score",
    "Application Status", "Deadline", "Apply Link", "Reason", "Evaluated At (UTC)", "Notes",
)
EXPECTED_REQUIREMENT_HEADERS = (
    "Tracker Row", "Job ID", "Requirement Type", "Requirement", "Candidate Value", "Status",
    "Confidence", "Method", "Reason Code", "Note",
)

#: Every leading character that must be neutralised (ADR-023 §6).
FORMULA_STARTS = ["=", "+", "-", "@", "\t", "\r"]


def requirement(**overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "requirement_type": "MIN_CGPA",
        "requirement": "Minimum CGPA 7.0 (SCALE_10)",
        "candidate_value": "8.2",
        "status": "PASS",
        "confidence": "HIGH",
        "method": "deterministic",
        "reason_code": "MEETS_MINIMUM",
        "note": "Meets the stated minimum.",
    }
    base.update(overrides)
    return base


def record(job_id: str = "job-1", **overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "job_id": job_id,
        "company_name": "Example Co",
        "role_title": "Software Engineering Intern",
    }
    base.update(overrides)
    return base


def full_record() -> dict[str, Any]:
    return record(
        "job-full",
        apply_link="https://example.com/careers/1",
        deadline="2026-11-15",
        job_status="ACTIVE",
        eligibility_state="LIKELY_ELIGIBLE",
        match_score=73.2,
        reason="Likely eligible. Shares Python and SQL.",
        application_status="INTERVIEW",
        evaluated_at="2026-09-19T15:30:00+05:30",
        notes="Second round booked.",
        requirement_breakdown=[
            requirement(),
            requirement(
                requirement_type="ALLOWED_FIELDS",
                requirement="Computer Science or related field",
                candidate_value="Information Technology",
                confidence="MEDIUM",
                method="ai_reasoning",
                reason_code="AI_FIELD_RELATED",
                note="Information Technology is a related field.",
            ),
        ],
    )


def parse(records: list[dict[str, Any]]) -> list[TrackerRecord]:
    return TrackerExportRequest(records=records).records


def export(records: list[dict[str, Any]]) -> bytes:
    return export_tracker(parse(records))


def workbook(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content))


def rows(sheet: Any) -> list[tuple[Any, ...]]:
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows()]


def zip_entries(content: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def content_of(content: bytes) -> list[tuple[str, list[list[tuple[Any, str, str]]]]]:
    """Every sheet's values, cell types and number formats — the file's meaning, not its bytes."""
    book = workbook(content)
    return [
        (
            sheet.title,
            [[(c.value, c.data_type, c.number_format) for c in row] for row in sheet.iter_rows()],
        )
        for sheet in book.worksheets
    ]


# ---------------------------------------------------------------------------------------
# Schema
# ---------------------------------------------------------------------------------------


def test_application_status_has_exactly_the_five_dossier_values() -> None:
    assert [status.value for status in ApplicationStatus] == [
        "NOT_APPLIED", "APPLIED", "INTERVIEW", "REJECTED", "OFFER",
    ]


def test_minimal_record_defaults() -> None:
    (parsed,) = parse([record()])
    assert parsed.application_status is ApplicationStatus.NOT_APPLIED
    assert parsed.match_score is None and parsed.requirement_breakdown is None


@pytest.mark.parametrize("count", [1, MAX_TRACKER_ROWS])
def test_row_count_bounds_accepted(count: int) -> None:
    assert len(parse([record(f"job-{i}") for i in range(count)])) == count


@pytest.mark.parametrize("count", [0, MAX_TRACKER_ROWS + 1])
def test_row_count_bounds_rejected(count: int) -> None:
    with pytest.raises(ValidationError):
        parse([record(f"job-{i}") for i in range(count)])


def test_row_cap_is_500() -> None:
    assert MAX_TRACKER_ROWS == 500


def test_duplicate_job_id_is_rejected() -> None:
    with pytest.raises(ValidationError, match="each job_id may appear only once"):
        parse([record("same"), record("other"), record("same")])


@pytest.mark.parametrize("score", [None, 0, 0.0, 55.5, 100])
def test_scores_accepted(score: float | None) -> None:
    assert parse([record(match_score=score)])[0].match_score == score


@pytest.mark.parametrize("score", [-0.1, 100.1, "high", float("nan"), float("inf")])
def test_invalid_scores_rejected(score: Any) -> None:
    with pytest.raises(ValidationError):
        parse([record(match_score=score)])


@pytest.mark.parametrize(
    "overrides",
    [
        {"application_status": "WITHDRAWN"},
        {"application_status": "applied"},
        {"eligibility_state": "MAYBE"},
        {"job_status": "OPEN"},
        {"deadline": "15/11/2026"},
        {"evaluated_at": "2026-09-19T10:00:00"},  # no time zone: rejected, not guessed
        {"unexpected": "value"},
        {"requirement_breakdown": [requirement(extra_field=1)]},
        {"requirement_breakdown": [requirement(status="MAYBE")]},
    ],
)
def test_invalid_record_rejected(overrides: dict[str, Any]) -> None:
    with pytest.raises(ValidationError):
        parse([record(**overrides)])


def test_unknown_top_level_field_rejected() -> None:
    with pytest.raises(ValidationError):
        TrackerExportRequest(records=[record()], candidate_id="nope")  # type: ignore[call-arg]


@pytest.mark.parametrize(
    ("field", "limit"),
    [("job_id", 100), ("company_name", 200), ("role_title", 200), ("apply_link", 2048),
     ("reason", 1000), ("notes", 1000)],
)
def test_length_limits(field: str, limit: int) -> None:
    parse([record(**{field: "x" * limit})])
    with pytest.raises(ValidationError):
        parse([record(**{field: "x" * (limit + 1)})])


def test_empty_job_id_rejected() -> None:
    with pytest.raises(ValidationError):
        parse([record("")])


@pytest.mark.parametrize("text_field", ["requirement", "candidate_value", "note"])
def test_requirement_text_is_bounded(text_field: str) -> None:
    parse([record(requirement_breakdown=[requirement(**{text_field: "x" * MAX_REQUIREMENT_TEXT})])])
    with pytest.raises(ValidationError):
        parse([record(requirement_breakdown=[
            requirement(**{text_field: "x" * (MAX_REQUIREMENT_TEXT + 1)})
        ])])


def test_requirement_count_is_bounded() -> None:
    parse([record(requirement_breakdown=[requirement()] * MAX_REQUIREMENTS_PER_ROW)])
    with pytest.raises(ValidationError):
        parse([record(requirement_breakdown=[requirement()] * (MAX_REQUIREMENTS_PER_ROW + 1))])


@pytest.mark.parametrize("char", ["\x00", "\x07", "\x0b", "\x1f"])
def test_unstorable_control_characters_rejected_without_echo(char: str) -> None:
    with pytest.raises(ValidationError) as caught:
        parse([record(notes=f"SECRET-NOTE{char}")])
    assert "SECRET-NOTE" not in str(caught.value.errors(include_input=False))
    with pytest.raises(ValidationError):
        parse([record(requirement_breakdown=[requirement(note=f"n{char}")])])


@pytest.mark.parametrize("char", ["\t", "\n", "\r"])
def test_tab_newline_and_return_are_storable(char: str) -> None:
    parse([record(notes=f"line one{char}line two")])


# ---------------------------------------------------------------------------------------
# Workbook structure
# ---------------------------------------------------------------------------------------


def test_header_order_is_fixed() -> None:
    assert TRACKER_HEADERS == EXPECTED_TRACKER_HEADERS
    assert REQUIREMENT_HEADERS == EXPECTED_REQUIREMENT_HEADERS
    book = workbook(export([full_record()]))
    assert rows(book[TRACKER_SHEET])[0] == EXPECTED_TRACKER_HEADERS
    assert rows(book[REQUIREMENTS_SHEET])[0] == EXPECTED_REQUIREMENT_HEADERS


def test_tracker_is_the_first_sheet_and_opens() -> None:
    book = workbook(export([record()]))
    assert book.sheetnames == [TRACKER_SHEET]
    assert book.active.title == TRACKER_SHEET


def test_every_field_is_exported_to_its_column() -> None:
    sheet = workbook(export([full_record()]))[TRACKER_SHEET]
    values = dict(zip(EXPECTED_TRACKER_HEADERS, rows(sheet)[1]))
    assert values == {
        "Row": 1,
        "Job ID": "job-full",
        "Company": "Example Co",
        "Role": "Software Engineering Intern",
        "Job Status": "ACTIVE",
        "Eligibility": "LIKELY_ELIGIBLE",
        "Match Score": 73.2,
        "Application Status": "INTERVIEW",
        "Deadline": datetime(2026, 11, 15),
        "Apply Link": "https://example.com/careers/1",
        "Reason": "Likely eligible. Shares Python and SQL.",
        "Evaluated At (UTC)": datetime(2026, 9, 19, 10, 0),
        "Notes": "Second round booked.",
    }


def test_optional_fields_absent_are_empty_cells() -> None:
    sheet = workbook(export([record()]))[TRACKER_SHEET]
    values = dict(zip(EXPECTED_TRACKER_HEADERS, rows(sheet)[1]))
    for header in ("Job Status", "Eligibility", "Match Score", "Deadline", "Apply Link",
                   "Reason", "Evaluated At (UTC)", "Notes"):
        assert values[header] is None, header
    assert values["Application Status"] == "NOT_APPLIED"


def test_row_order_is_preserved_exactly() -> None:
    ids = ["zeta", "alpha", "mid", "0-first", "beta"]
    sheet = workbook(export([record(job_id) for job_id in ids]))[TRACKER_SHEET]
    body = rows(sheet)[1:]
    assert [row[1] for row in body] == ids
    assert [row[0] for row in body] == [1, 2, 3, 4, 5]


def test_one_row_workbook() -> None:
    sheet = workbook(export([record()]))[TRACKER_SHEET]
    assert sheet.max_row == 2


def test_five_hundred_row_workbook() -> None:
    records = [record(f"job-{i:03d}", match_score=i / 5) for i in range(MAX_TRACKER_ROWS)]
    sheet = workbook(export(records))[TRACKER_SHEET]
    body = rows(sheet)[1:]
    assert len(body) == MAX_TRACKER_ROWS
    assert body[0][:2] == (1, "job-000") and body[-1][:2] == (500, "job-499")
    assert [row[6] for row in body] == [i / 5 for i in range(MAX_TRACKER_ROWS)]


def test_every_application_status_is_written() -> None:
    records = [record(status.value, application_status=status.value) for status in ApplicationStatus]
    body = rows(workbook(export(records))[TRACKER_SHEET])[1:]
    assert [row[7] for row in body] == [status.value for status in ApplicationStatus]


def test_unicode_survives() -> None:
    text = "Zürich — 日本語 · café 🚀"
    sheet = workbook(export([record(company_name=text, notes=text, reason=text)]))[TRACKER_SHEET]
    values = dict(zip(EXPECTED_TRACKER_HEADERS, rows(sheet)[1]))
    assert values["Company"] == values["Notes"] == values["Reason"] == text


# ---------------------------------------------------------------------------------------
# Scores and dates
# ---------------------------------------------------------------------------------------


def test_null_score_is_an_empty_cell_not_zero() -> None:
    sheet = workbook(export([
        record("scored", match_score=0.0),
        record("withheld", match_score=None, eligibility_state="NOT_ELIGIBLE",
               reason="Not eligible: minimum CGPA not met. Score withheld."),
    ]))[TRACKER_SHEET]
    zero, null = sheet["G2"], sheet["G3"]
    assert zero.value == 0 and zero.data_type == "n"
    assert null.value is None
    assert sheet["K3"].value == "Not eligible: minimum CGPA not met. Score withheld."


def test_null_score_writes_no_value_in_the_xml() -> None:
    sheet_xml = zip_entries(export([record(match_score=None)]))["xl/worksheets/sheet1.xml"]
    assert b'r="G2"' not in sheet_xml  # no cell at all, not a 0


def test_score_is_numeric() -> None:
    cell = workbook(export([record(match_score=42.5)]))[TRACKER_SHEET]["G2"]
    assert cell.value == 42.5 and cell.data_type == "n"


def test_deadline_is_a_real_date() -> None:
    cell = workbook(export([record(deadline="2026-11-15")]))[TRACKER_SHEET]["I2"]
    assert cell.is_date and cell.data_type == "d"
    assert cell.value.date() == date(2026, 11, 15)
    assert cell.number_format == "yyyy-mm-dd"


def test_evaluated_at_is_a_utc_datetime() -> None:
    cell = workbook(export([record(evaluated_at="2026-09-19T23:30:00-04:00")]))[TRACKER_SHEET]["L2"]
    assert cell.is_date
    assert cell.value == datetime(2026, 9, 20, 3, 30)
    assert cell.number_format == "yyyy-mm-dd hh:mm:ss"


def test_evaluated_at_in_utc_is_unchanged() -> None:
    parsed = parse([record(evaluated_at=datetime(2026, 1, 2, 3, 4, 5, tzinfo=timezone.utc))])
    cell = workbook(export_tracker(parsed))[TRACKER_SHEET]["L2"]
    assert cell.value == datetime(2026, 1, 2, 3, 4, 5)


# ---------------------------------------------------------------------------------------
# Requirements sheet
# ---------------------------------------------------------------------------------------


def test_requirements_sheet_lists_each_requirement_with_its_parent_row() -> None:
    records = [
        record("no-breakdown"),
        full_record(),
        record("one", requirement_breakdown=[requirement(status="FAIL", reason_code="BELOW_MINIMUM")]),
    ]
    book = workbook(export(records))
    assert book.sheetnames == [TRACKER_SHEET, REQUIREMENTS_SHEET]
    body = rows(book[REQUIREMENTS_SHEET])[1:]
    assert body == [
        (2, "job-full", "MIN_CGPA", "Minimum CGPA 7.0 (SCALE_10)", "8.2", "PASS", "HIGH",
         "deterministic", "MEETS_MINIMUM", "Meets the stated minimum."),
        (2, "job-full", "ALLOWED_FIELDS", "Computer Science or related field",
         "Information Technology", "PASS", "MEDIUM", "ai_reasoning", "AI_FIELD_RELATED",
         "Information Technology is a related field."),
        (3, "one", "MIN_CGPA", "Minimum CGPA 7.0 (SCALE_10)", "8.2", "FAIL", "HIGH",
         "deterministic", "BELOW_MINIMUM", "Meets the stated minimum."),
    ]
    tracker = rows(book[TRACKER_SHEET])
    for parent_row, job_id, *_ in body:
        assert tracker[parent_row][:2] == (parent_row, job_id)


def test_requirement_without_candidate_value_is_empty() -> None:
    book = workbook(export([record(requirement_breakdown=[requirement(candidate_value=None)])]))
    assert book[REQUIREMENTS_SHEET]["E2"].value is None


@pytest.mark.parametrize("breakdown", [None, []])
def test_requirements_sheet_absent_without_breakdown_data(breakdown: list | None) -> None:
    book = workbook(export([record(requirement_breakdown=breakdown), record("b")]))
    assert book.sheetnames == [TRACKER_SHEET]


# ---------------------------------------------------------------------------------------
# Formula injection
# ---------------------------------------------------------------------------------------


@pytest.mark.parametrize("start", FORMULA_STARTS)
def test_neutralize_prefixes_every_formula_start(start: str) -> None:
    assert neutralize(f"{start}SUM(A1:A2)") == f"'{start}SUM(A1:A2)"


@pytest.mark.parametrize("text", ["Example Co", "", "a=b", " =1", "'already", "1+1", "#REF"])
def test_neutralize_leaves_safe_text_alone(text: str) -> None:
    assert neutralize(text) == text


TEXT_COLUMNS = {"job_id": "B", "company_name": "C", "role_title": "D", "apply_link": "J",
                "reason": "K", "notes": "M"}


@pytest.mark.parametrize("start", FORMULA_STARTS)
@pytest.mark.parametrize("field", list(TEXT_COLUMNS))
def test_every_text_field_is_literal_text(field: str, start: str) -> None:
    payload = f"{start}HYPERLINK(\"http://attacker.invalid\",\"x\")"
    sheet = workbook(export([record(**{field: payload})]))[TRACKER_SHEET]
    cell = sheet[f"{TEXT_COLUMNS[field]}2"]
    assert cell.data_type == "s"
    assert cell.value == "'" + payload


@pytest.mark.parametrize("start", FORMULA_STARTS)
@pytest.mark.parametrize("field", ["requirement", "candidate_value", "note"])
def test_requirement_text_is_literal_text(field: str, start: str) -> None:
    payload = f"{start}1+1"
    book = workbook(export([record(requirement_breakdown=[requirement(**{field: payload})])]))
    column = {"requirement": "D", "candidate_value": "E", "note": "J"}[field]
    cell = book[REQUIREMENTS_SHEET][f"{column}2"]
    assert cell.data_type == "s" and cell.value == "'" + payload


def test_no_formula_element_in_any_sheet_xml() -> None:
    hostile = [
        record(f"{start}id-{i}", company_name=f"{start}cmd|' /C calc'!A0",
               role_title=f"{start}1+1", apply_link=f"{start}HYPERLINK(\"x\")",
               reason=f"{start}SUM(1)", notes=f"{start}2*3",
               requirement_breakdown=[requirement(requirement=f"{start}A1", note=f"{start}B1",
                                                  candidate_value=f"{start}C1")])
        for i, start in enumerate(FORMULA_STARTS)
    ]
    entries = zip_entries(export(hostile))
    sheets = [name for name in entries if name.startswith("xl/worksheets/sheet")]
    assert len(sheets) == 2
    for name in sheets:
        assert b"<f>" not in entries[name] and b"<f " not in entries[name], name
    assert "xl/calcChain.xml" not in entries


@pytest.mark.parametrize("code", ["#N/A", "#REF!", "#DIV/0!", "#NULL!", "#VALUE!", "#NAME?", "#NUM!"])
def test_error_code_text_stays_text(code: str) -> None:
    """openpyxl would store these strings as Excel error cells; the pinned type keeps them text."""
    sheet = workbook(export([record(notes=code, reason=code)]))[TRACKER_SHEET]
    for cell in (sheet["K2"], sheet["M2"]):
        assert cell.data_type == "s" and cell.value == code


def test_formula_prefix_alone_would_have_become_a_formula() -> None:
    """Guard the guard: without neutralisation openpyxl really does store a formula."""
    book = openpyxl.Workbook()
    book.active["A1"] = "=1+1"
    assert book.active["A1"].data_type == "f"


def test_links_are_plain_text_not_hyperlinks() -> None:
    content = export([record(apply_link="https://example.com/careers/1")])
    sheet = workbook(content)[TRACKER_SHEET]
    assert sheet["J2"].hyperlink is None and sheet["J2"].value == "https://example.com/careers/1"
    entries = zip_entries(content)
    assert not any("hyperlink" in name.lower() for name in entries)
    assert not any(name.startswith("xl/worksheets/_rels/") for name in entries)
    for name, data in entries.items():
        assert b"<hyperlink" not in data, name


# ---------------------------------------------------------------------------------------
# Metadata, identity, determinism
# ---------------------------------------------------------------------------------------

NOTE_MARKER = "NOTE-MARKER-6a1f"
COMPANY_MARKER = "COMPANY-MARKER-6a1f"


def test_metadata_is_fixed_and_carries_no_row_data() -> None:
    content = export([record(COMPANY_MARKER, company_name=COMPANY_MARKER, notes=NOTE_MARKER)])
    properties = workbook(content).properties
    assert properties.creator == WORKBOOK_CREATOR == "EligiCore"
    assert properties.lastModifiedBy == "EligiCore"
    assert properties.title == WORKBOOK_TITLE
    for field in ("subject", "description", "keywords", "category", "identifier", "language"):
        assert getattr(properties, field) in (None, ""), field
    entries = zip_entries(content)
    for name in ("docProps/core.xml", "docProps/app.xml", "xl/workbook.xml"):
        for marker in (NOTE_MARKER, COMPANY_MARKER):
            assert marker.encode() not in entries[name], (name, marker)


def test_no_comments_hidden_sheets_or_defined_names() -> None:
    content = export([full_record()])
    book = workbook(content)
    assert not book.defined_names
    for sheet in book.worksheets:
        assert sheet.sheet_state == "visible"
        assert not any(dim.hidden for dim in sheet.column_dimensions.values())
        assert not any(dim.hidden for dim in sheet.row_dimensions.values())
        assert not any(cell.comment for row in sheet.iter_rows() for cell in row)
    assert not any("comment" in name.lower() or "vba" in name.lower() for name in zip_entries(content))


def test_filename_is_fixed() -> None:
    assert TRACKER_FILENAME == "eligicore-tracker.xlsx"


def test_same_records_give_the_same_content() -> None:
    records = [full_record(), record("b", notes="x"), record("c", match_score=None)]
    first, second = export(records), export(records)
    assert content_of(first) == content_of(second)
    # Only the document timestamps may differ between runs.
    a, b = zip_entries(first), zip_entries(second)
    assert a.keys() == b.keys()
    assert {name for name in a if a[name] != b[name]} <= {"docProps/core.xml"}


# ---------------------------------------------------------------------------------------
# In memory only
# ---------------------------------------------------------------------------------------


def test_in_memory_writer_matches_openpyxls_own_save() -> None:
    """The override changes where worksheets are buffered, nothing in the file itself."""
    records = parse([full_record(), record("b", notes="=x"), record("c", match_score=None)])
    ours = zip_entries(export_tracker(records))
    stock_buffer = io.BytesIO()
    build_workbook(records).save(stock_buffer)
    stock = zip_entries(stock_buffer.getvalue())
    assert list(ours) == list(stock)
    assert {name for name in ours if ours[name] != stock[name]} <= {"docProps/core.xml"}


def test_export_opens_no_file(monkeypatch: pytest.MonkeyPatch) -> None:
    content = export([record()])  # warm imports first

    def refuse(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("the export touched the filesystem")

    for target, name in [(builtins, "open"), (io, "open"), (os, "open"),
                         (tempfile, "mkstemp"), (tempfile, "mkdtemp"),
                         (tempfile, "NamedTemporaryFile"), (tempfile, "TemporaryFile"),
                         (tempfile, "SpooledTemporaryFile")]:
        monkeypatch.setattr(target, name, refuse)
    assert content_of(export([record()])) == content_of(content)


def test_export_writes_nothing_to_cwd_or_temp(tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch) -> None:
    work, temp = tmp_path / "work", tmp_path / "temp"
    work.mkdir()
    temp.mkdir()
    monkeypatch.chdir(work)
    monkeypatch.setattr(tempfile, "tempdir", str(temp))
    export([full_record()] + [record(f"r{i}") for i in range(50)])
    assert os.listdir(work) == [] and os.listdir(temp) == []


# ---------------------------------------------------------------------------------------
# Architecture
# ---------------------------------------------------------------------------------------


def imported_modules(path: pathlib.Path) -> set[str]:
    modules: set[str] = set()
    for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
        if isinstance(node, ast.Import):
            modules.update(alias.name for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            modules.add(node.module)
    return modules


FORBIDDEN_PREFIXES = (
    "fastapi", "starlette", "sqlalchemy", "httpx", "app.ai", "app.models", "app.database",
    "app.routers", "app.main", "app.services.eligibility_engine", "app.services.eligibility_ai",
    "app.services.matching_engine", "app.services.recommendations", "app.services.job_ingestion",
    "os", "pathlib", "tempfile", "shutil", "sqlite3", "pickle", "glob",
)


def test_service_imports_no_framework_database_ai_engine_or_filesystem() -> None:
    for module in imported_modules(SERVICE_PATH):
        for forbidden in FORBIDDEN_PREFIXES:
            assert module != forbidden and not module.startswith(forbidden + "."), module


def test_service_uses_io_only_for_an_in_memory_buffer() -> None:
    tree = ast.parse(SERVICE_PATH.read_text(encoding="utf-8"))
    io_names = {
        alias.name
        for node in ast.walk(tree)
        if isinstance(node, ast.ImportFrom) and node.module == "io"
        for alias in node.names
    }
    assert io_names == {"BytesIO"}
    assert "io" not in {alias.name for node in ast.walk(tree) if isinstance(node, ast.Import)
                        for alias in node.names}
    calls = {node.func.id for node in ast.walk(tree)
             if isinstance(node, ast.Call) and isinstance(node.func, ast.Name)}
    assert "open" not in calls


def test_service_depends_only_on_openpyxl_schemas_and_stdlib() -> None:
    for module in imported_modules(SERVICE_PATH):
        assert module.split(".")[0] in {
            "__future__", "logging", "time", "collections", "datetime", "enum", "io", "zipfile",
            "openpyxl",
            "app",
        }, module
        if module.startswith("app."):
            assert module.startswith("app.schemas."), module


@pytest.mark.parametrize(
    "path",
    sorted(
        str(p) for p in pathlib.Path("app").rglob("*.py")
        if p.name not in {"tracker_export.py", "applications.py", "tracker.py", "main.py"}
    ),
)
def test_nothing_else_imports_the_export(path: str) -> None:
    for module in imported_modules(pathlib.Path(path)):
        assert "tracker" not in module and "applications" not in module, (path, module)


def test_export_takes_no_database_or_ai_dependency() -> None:
    """Asserted on the handler, because the module now holds a second endpoint.

    ``/applications/prepare`` legitimately takes a session and an AI service (ADR-025), so a
    module-level import check would pass for the wrong reason. The guarantee that matters is
    per-endpoint: the export reads nothing and calls no provider, so its signature has one
    parameter — the request — and no dependency at all.
    """
    import inspect

    from app.routers.applications import export

    assert set(inspect.signature(export).parameters) == {"request"}

    # Whatever else lands in this module, the export's neighbours stay out of it: no endpoint
    # here may reach eligibility, matching or recommendations (ADR-023 §11, ADR-025 D7).
    modules = imported_modules(pathlib.Path("app/routers/applications.py"))
    assert not any(m.startswith(("app.services.recommendations", "app.services.matching_engine",
                                 "app.services.eligibility")) for m in modules), modules


def test_service_is_importable_without_the_application() -> None:
    import subprocess
    import sys

    code = (
        "import sys; import app.services.tracker_export as t; "
        "assert 'fastapi' not in sys.modules and 'sqlalchemy' not in sys.modules, "
        "sorted(m for m in sys.modules if m.startswith(('fastapi', 'sqlalchemy')))"
    )
    result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
