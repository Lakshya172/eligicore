"""Builders for synthetic resume documents used in tests.

Real PDFs and DOCX files, generated in-memory. Nothing is read from disk and **no real
resume is ever committed to this repository** (``standards/testing.md`` §5, QG-005).

The PDF is assembled by hand rather than with a rendering library. Adding reportlab just to
produce test fixtures would put a dependency in the tree that ships nothing — the resulting
file is a genuine PDF that pdfplumber parses normally.

Every value produced here is obviously fictional.
"""

from __future__ import annotations

import io

from docx import Document

# A synthetic resume. Deliberately fictional: "Test Candidate", example.com, an all-zero
# phone number, and invented institutions.
SAMPLE_RESUME_LINES: tuple[str, ...] = (
    "Test Candidate",
    "test.candidate@example.com",
    "+10000000000",
    "Example City",
    "",
    "EDUCATION",
    "B.Tech Information Technology",
    "Example Institute of Technology",
    "CGPA 8.2/10",
    "Graduating 2027",
    "",
    "SKILLS",
    "Python, React, SQL, Docker, Git",
    "",
    "EXPERIENCE",
    "Software Engineering Intern at Example Corp",
    "Jun 2026 - Aug 2026",
)

# Same resume, but the grade carries no scale. Used to prove an unstated scale stays
# UNKNOWN rather than being assumed to be out of 10.
SAMPLE_RESUME_NO_SCALE_LINES: tuple[str, ...] = (
    "Test Candidate",
    "test.candidate@example.com",
    "",
    "EDUCATION",
    "B.Tech Information Technology",
    "CGPA 8.2",
    "Graduating 2027",
    "",
    "SKILLS",
    "Python, React",
)


def build_pdf(lines: tuple[str, ...] | list[str]) -> bytes:
    """Build a minimal single-page PDF containing ``lines`` as visible text."""
    content = "BT\n/F1 11 Tf\n50 750 Td\n14 TL\n"
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content += f"({escaped}) Tj\nT*\n"
    content += "ET"
    stream = content.encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")

    xref_offset = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return out.getvalue()


def build_docx(lines: tuple[str, ...] | list[str]) -> bytes:
    """Build a DOCX containing ``lines`` as paragraphs."""
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_docx_with_table(rows: list[tuple[str, str]]) -> bytes:
    """Build a DOCX whose content lives in a table.

    Resumes frequently lay education and skills out in tables, and reading paragraphs alone
    silently misses all of it.
    """
    document = Document()
    document.add_paragraph("Test Candidate")
    table = document.add_table(rows=len(rows), cols=2)
    for row_index, (left, right) in enumerate(rows):
        table.rows[row_index].cells[0].text = left
        table.rows[row_index].cells[1].text = right
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_empty_pdf() -> bytes:
    """A structurally valid PDF page with no text on it — e.g. a scanned image."""
    return build_pdf([])


def build_corrupt_pdf() -> bytes:
    """Bytes that begin like a PDF but are not one."""
    return b"%PDF-1.4\nthis is not a valid pdf body at all\n%%EOF\n"


def build_corrupt_docx() -> bytes:
    """Bytes that begin like a ZIP container but are not a valid DOCX."""
    return b"PK\x03\x04" + b"\x00" * 64
