"""The Phase 1 flow, end to end and offline (PR 6B — ADR-024 §5).

One test file walks the dossier's §7 journey across the real HTTP API: résumé → parsed profile →
the candidate completing the gaps (step 3) → validation and normalization → eligibility →
recommendations → Excel tracker. It uses the seeded 40-job catalogue and the mock AI provider, and
a socket guard fails the run if anything reaches for the network.

Two distinct synthetic profiles are exercised, because §17 asks for recommendations that rank
sensibly *across* profiles: an IT graduate and a mechanical engineer see different catalogues.

Group expectations are golden, keyed by curated `source_job_id`. A catalogue edit that changes an
outcome must update them in the same change — that is the point.
"""

from __future__ import annotations

import io
import json
import logging
import os
import pathlib
import socket
import tempfile
from collections.abc import Iterator
from dataclasses import dataclass
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.ai.ai_service import AIService, get_lazy_ai_service
from app.ai.providers.mock import MockAIProvider
from app.cli import seed_catalogue
from app.database import Base, get_db
from app.main import app
from app.schemas.candidate import Confidence
from app.schemas.eligibility import FieldRelatedness, FieldRelatednessAssessment
from tests.conftest import ALLOWED_OPERATIONAL_TABLES
from tests.fixtures_documents import SAMPLE_RESUME_LINES, build_pdf

# Values the synthetic résumé fixture carries. Nothing about them may reach the database, a log
# line or an error response.
MARKERS = (
    "Test Candidate",
    "test.candidate@example.com",
    "+10000000000",
    "Example Institute of Technology",
)

# --- golden outcomes over the seeded catalogue (ADR-024 §2, §5) -------------------------

A_RANKED = {
    "EX-FT-040", "EX-FT-034", "EX-INT-029", "EX-FT-010", "EX-INT-001", "EX-FT-007", "EX-INT-005",
    "EX-INT-011", "EX-INT-017", "EX-INT-006", "EX-FT-018", "EX-FT-026", "EX-FT-020", "EX-FT-012",
    "EX-INT-033", "EX-FT-036", "EX-INT-031", "EX-INT-023", "EX-INT-039",
}
A_NEEDS_REVIEW = {
    "EX-INT-025", "EX-INT-035", "EX-INT-008", "EX-INT-019", "EX-FT-016", "EX-INT-003",
    "EX-INT-027", "EX-INT-021", "EX-INT-037", "EX-INT-013", "EX-INT-015", "EX-FT-038",
}
A_NOT_ELIGIBLE = {
    "EX-FT-004", "EX-FT-022", "EX-FT-028", "EX-FT-032", "EX-FT-014", "EX-FT-030", "EX-INT-009",
    "EX-FT-002", "EX-FT-024",
}
B_RANKED = {"EX-FT-014", "EX-FT-030", "EX-INT-031", "EX-FT-026", "EX-FT-007", "EX-INT-039"}
B_NEEDS_REVIEW = {"EX-FT-002", "EX-INT-003", "EX-INT-023", "EX-FT-036", "EX-FT-016", "EX-FT-038"}
#: Jobs whose permitted fields the mock only accepts once relatedness is pinned to RELATED.
A_LIKELY_WHEN_RELATED = {
    "EX-INT-035", "EX-INT-008", "EX-FT-016", "EX-INT-027", "EX-INT-021", "EX-INT-013",
}
#: The general-application listing: no description, no skills, stop-word title (`NO_JOB_TERMS`).
NO_TERMS_JOB = "EX-INT-039"


@dataclass
class Flow:
    client: TestClient
    engine: Any
    ids: dict[str, str]  # catalogue job id -> curated source_job_id


LOOPBACK = {"127.0.0.1", "::1", "localhost", "0.0.0.0"}


def _is_loopback(address: Any) -> bool:
    """Loopback stays allowed: asyncio's event loop uses a self-pipe over sockets on Windows."""
    host = address[0] if isinstance(address, tuple) and address else address
    return isinstance(host, str) and host in LOOPBACK


@pytest.fixture(autouse=True)
def no_network(monkeypatch: pytest.MonkeyPatch) -> None:
    """Fail the test if anything reaches off the machine. The suite runs offline (memory I-1)."""
    real_connect = socket.socket.connect
    real_connect_ex = socket.socket.connect_ex
    real_create = socket.create_connection

    def guard(original: Any) -> Any:
        def wrapper(self_or_address: Any, *args: Any, **kwargs: Any) -> Any:
            address = args[0] if args else self_or_address
            if not _is_loopback(address):
                raise AssertionError(f"the full flow attempted a network connection: {address!r}")
            return original(self_or_address, *args, **kwargs)

        return wrapper

    monkeypatch.setattr(socket.socket, "connect", guard(real_connect))
    monkeypatch.setattr(socket.socket, "connect_ex", guard(real_connect_ex))
    monkeypatch.setattr(socket, "create_connection", guard(real_create))


@pytest.fixture
def flow() -> Iterator[Flow]:
    """The real application over an in-memory database seeded by the demo path."""
    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as session:
        outcomes = seed_catalogue(session)  # the same function `python -m app.cli` runs
    assert [o.jobs_created for o in outcomes] == [40]

    def override_get_db() -> Iterator[Session]:
        db = factory()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    ids: dict[str, str] = {}
    offset = 0
    while True:
        page = client.get(f"/api/v1/jobs?limit=100&offset={offset}").json()
        ids.update({job["id"]: job["source_job_id"] for job in page["items"]})
        offset += page["limit"]
        if offset >= page["total"]:
            break
    assert len(ids) == 40
    try:
        yield Flow(client, engine, ids)
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_lazy_ai_service, None)
        engine.dispose()


# ---------------------------------------------------------------------------------------
# Profiles
# ---------------------------------------------------------------------------------------


def parse_resume(flow: Flow) -> dict[str, Any]:
    response = flow.client.post(
        "/api/v1/resumes/parse",
        files={"file": ("resume.pdf", build_pdf(SAMPLE_RESUME_LINES), "application/pdf")},
    )
    assert response.status_code == 200
    return response.json()


def complete_profile(parsed: dict[str, Any]) -> dict[str, Any]:
    """Dossier §7 step 3: the candidate fills what parsing could not determine."""
    profile = parsed["profile"]
    education = dict(profile["education"][0])
    education.update(field_of_study="Information Technology", level="BACHELORS")
    return {**profile, "education": [education], "backlogs": 0}


PROFILE_B: dict[str, Any] = {
    "candidate_id": "profile-b-synthetic",
    "education": [
        {
            "degree": "B.E.",
            "level": "BACHELORS",
            "field_of_study": "Mechanical Engineering",
            "grad_year": 2026,
            "cgpa": 3.2,
            "scale": "SCALE_4",
        }
    ],
    "backlogs": 1,
    "skills": ["SolidWorks", "AutoCAD", "MATLAB", "Python"],
    "experience": [
        {
            "title": "Design Intern",
            "description": "CAD modelling of fixtures in SolidWorks and tolerance studies in MATLAB.",
        }
    ],
}


def recommend(flow: Flow, profile: dict[str, Any]) -> dict[str, Any]:
    response = flow.client.post("/api/v1/recommendations", json={"profile": profile})
    assert response.status_code == 200
    return response.json()


def grouped_ids(flow: Flow, body: dict[str, Any], group: str) -> list[str]:
    """Curated ids in response order."""
    return [flow.ids[item["eligibility"]["job_id"]] for item in body[group]]


def grouped(flow: Flow, body: dict[str, Any], group: str) -> set[str]:
    """Curated ids as a set.

    Group *membership* is deterministic; the order within a group is not always, because a
    catalogue id is a UUID minted at seed time and both the unranked groups and score ties
    order by it. Ordering is asserted separately, on scores and ranks.
    """
    return set(grouped_ids(flow, body, group))


def tracker_records(flow: Flow, body: dict[str, Any]) -> list[dict[str, Any]]:
    """Build tracker rows from a recommendation response, as a client would."""
    records = []
    for group in ("ranked", "needs_review", "not_eligible"):
        for item in body[group]:
            eligibility = item["eligibility"]
            records.append(
                {
                    "job_id": eligibility["job_id"],
                    "company_name": eligibility["company_name"],
                    "role_title": eligibility["role_title"],
                    "apply_link": eligibility["apply_link"],
                    "deadline": eligibility["deadline"],
                    "job_status": eligibility["job_status"],
                    "eligibility_state": eligibility["eligibility_state"],
                    "match_score": item["match"]["match_score"],
                    "reason": item["explanation"][:1000],
                    "application_status": "APPLIED" if group == "ranked" else "NOT_APPLIED",
                    "evaluated_at": body["evaluated_at"],
                    "requirement_breakdown": eligibility["requirement_breakdown"][:20] or None,
                }
            )
    return records


# ---------------------------------------------------------------------------------------
# Profile A — the whole journey
# ---------------------------------------------------------------------------------------


def test_profile_a_walks_the_full_flow(flow: Flow) -> None:
    # 1–2. Résumé upload and parsing.
    parsed = parse_resume(flow)
    assert parsed["status"] in {"PARSED", "PARTIAL"}
    assert parsed["profile"]["skills"]
    assert parsed["profile"]["education"][0]["cgpa"] == 8.2
    assert parsed["profile"]["education"][0]["scale"] == "SCALE_10"

    # 3. Profile completion is genuinely required: parsing leaves these unknown.
    assert parsed["profile"]["education"][0]["field_of_study"] is None
    assert parsed["profile"]["education"][0]["level"] == "UNKNOWN"
    assert parsed["profile"]["backlogs"] is None
    assert "backlogs" in {issue["field"] for issue in parsed["issues"]}
    profile = complete_profile(parsed)

    # 4. Validation and normalization.
    validated = flow.client.post("/api/v1/candidates/validate", json=profile)
    assert validated.status_code == 200 and validated.json()["is_valid"] is True
    normalized = flow.client.post("/api/v1/candidates/normalize", json=profile)
    assert normalized.status_code == 200
    assert "Python" in normalized.json()["profile"]["skills"]

    # 5. Eligibility for a representative subset.
    subset = {"EX-FT-040": "ELIGIBLE", "EX-FT-002": "NOT_ELIGIBLE", "EX-INT-003": "NEEDS_REVIEW"}
    by_source = {source: job_id for job_id, source in flow.ids.items()}
    checked = flow.client.post(
        "/api/v1/eligibility/check",
        json={"profile": profile, "job_ids": [by_source[s] for s in subset]},
    )
    assert checked.status_code == 200
    verdicts = {
        flow.ids[r["job_id"]]: r["eligibility_state"] for r in checked.json()["results"]
    }
    assert verdicts == subset
    for result in checked.json()["results"]:
        assert result["requirement_breakdown"] and result["summary"]

    # 6. Recommendations over the whole catalogue.
    body = recommend(flow, profile)
    assert body["jobs_considered"] == 40 and body["jobs_not_considered"] == 0
    assert grouped(flow, body, "ranked") == A_RANKED
    assert grouped(flow, body, "needs_review") == A_NEEDS_REVIEW
    assert grouped(flow, body, "not_eligible") == A_NOT_ELIGIBLE
    assert grouped_ids(flow, body, "ranked")[0] == "EX-FT-040", "highest score leads"
    assert grouped_ids(flow, body, "ranked")[-1] == NO_TERMS_JOB, "the unscored job sorts last"
    assert body["not_open"] == [] and body["not_found_job_ids"] == []

    # Ranking: descending score, nulls last, ranks 1..n.
    for group in ("ranked", "needs_review"):
        scores = [item["match"]["match_score"] for item in body[group]]
        present = [s for s in scores if s is not None]
        assert present == sorted(present, reverse=True)
        assert all(s is None for s in scores[len(present):])
        assert [item["rank"] for item in body[group]] == list(range(1, len(scores) + 1))

    # The eligibility verdicts inside recommendations are the same ones /eligibility/check gives.
    embedded = {
        flow.ids[item["eligibility"]["job_id"]]: item["eligibility"]["eligibility_state"]
        for group in ("ranked", "needs_review", "not_eligible")
        for item in body[group]
    }
    assert {s: embedded[s] for s in subset} == subset

    # Withheld scores, and the job with no matching terms.
    for item in body["not_eligible"]:
        assert item["match"]["match_score"] is None
        assert item["match"]["score_basis"] == "WITHHELD_NOT_ELIGIBLE"
        assert item["rank"] is None
    no_terms = [i for i in body["ranked"] if flow.ids[i["eligibility"]["job_id"]] == NO_TERMS_JOB]
    assert no_terms and no_terms[0]["match"]["match_score"] is None
    assert no_terms[0]["match"]["score_basis"] == "NO_JOB_TERMS"

    # 7. Excel tracker export.
    records = tracker_records(flow, body)
    assert len(records) == 40
    export = flow.client.post("/api/v1/applications/export", json={"records": records})
    assert export.status_code == 200
    assert export.headers["content-disposition"] == 'attachment; filename="eligicore-tracker.xlsx"'

    book = openpyxl.load_workbook(io.BytesIO(export.content))
    assert book.sheetnames == ["Tracker", "Requirements"]
    tracker = book["Tracker"]
    assert tracker.max_row == 41
    exported = [tracker.cell(row=r, column=2).value for r in range(2, 42)]
    assert exported == [record["job_id"] for record in records], "row order preserved"
    states = [tracker.cell(row=r, column=6).value for r in range(2, 42)]
    scores = [tracker.cell(row=r, column=7).value for r in range(2, 42)]
    assert states.count("NOT_ELIGIBLE") == len(A_NOT_ELIGIBLE)
    for state, score in zip(states, scores):
        if state == "NOT_ELIGIBLE":
            assert score is None, "a withheld score is an empty cell, never 0"
    assert book["Requirements"].max_row > 1


# ---------------------------------------------------------------------------------------
# Profile B — a different candidate sees a different catalogue
# ---------------------------------------------------------------------------------------


def test_profile_b_flows_through_to_an_export(flow: Flow) -> None:
    validated = flow.client.post("/api/v1/candidates/validate", json=PROFILE_B)
    assert validated.status_code == 200 and validated.json()["is_valid"] is True

    body = recommend(flow, PROFILE_B)
    assert grouped(flow, body, "ranked") == B_RANKED
    assert grouped(flow, body, "needs_review") == B_NEEDS_REVIEW
    assert len(body["not_eligible"]) == 28

    # The mechanical roles lead, which is the point of matching.
    assert grouped_ids(flow, body, "ranked")[:2] == ["EX-FT-014", "EX-FT-030"]  # distinct scores
    top = body["ranked"][0]
    assert top["match"]["match_score"] > body["ranked"][1]["match"]["match_score"]
    assert {t["term"] for t in top["match"]["top_terms"]} & {"SolidWorks", "AutoCAD", "MATLAB"}

    export = flow.client.post(
        "/api/v1/applications/export", json={"records": tracker_records(flow, body)}
    )
    assert export.status_code == 200
    assert openpyxl.load_workbook(io.BytesIO(export.content))["Tracker"].max_row == 41


def test_the_two_profiles_get_different_recommendations(flow: Flow) -> None:
    a = recommend(flow, complete_profile(parse_resume(flow)))
    b = recommend(flow, PROFILE_B)

    a_ranked, b_ranked = grouped_ids(flow, a, "ranked"), grouped_ids(flow, b, "ranked")
    assert set(a_ranked) != set(b_ranked)
    assert a_ranked[0] != b_ranked[0]
    assert len(a["not_eligible"]) != len(b["not_eligible"])
    # Each profile's top job is ineligible for the other.
    assert b_ranked[0] in grouped(flow, a, "not_eligible")
    assert a_ranked[0] in grouped(flow, b, "not_eligible")
    assert a["corpus_fingerprint"] == b["corpus_fingerprint"], "same catalogue, same corpus"


# ---------------------------------------------------------------------------------------
# AI: mock only, and the LIKELY_ELIGIBLE path
# ---------------------------------------------------------------------------------------


def test_default_mock_keeps_ambiguous_fields_in_needs_review(flow: Flow) -> None:
    """The runtime default mock answers UNCERTAIN, so nothing is promoted (ADR-019)."""
    body = recommend(flow, complete_profile(parse_resume(flow)))
    assert A_LIKELY_WHEN_RELATED <= grouped(flow, body, "needs_review")
    states = {item["eligibility"]["eligibility_state"] for item in body["ranked"]}
    assert "LIKELY_ELIGIBLE" not in states


def test_pinned_related_answer_produces_likely_eligible(flow: Flow) -> None:
    provider = MockAIProvider(
        return_relatedness=FieldRelatednessAssessment(
            result=FieldRelatedness.RELATED,
            confidence=Confidence.MEDIUM,
            reason="Synthetic pinned answer for the demo flow.",
        )
    )
    app.dependency_overrides[get_lazy_ai_service] = lambda: AIService(provider)

    body = recommend(flow, complete_profile(parse_resume(flow)))

    likely = [
        flow.ids[item["eligibility"]["job_id"]]
        for item in body["ranked"]
        if item["eligibility"]["eligibility_state"] == "LIKELY_ELIGIBLE"
    ]
    assert set(likely) == A_LIKELY_WHEN_RELATED
    assert provider.relatedness_call_count > 0, "the AI stage really ran"
    # AI can promote to LIKELY_ELIGIBLE but never rescue a deterministic failure (INV-2).
    assert grouped(flow, body, "not_eligible") == A_NOT_ELIGIBLE


# ---------------------------------------------------------------------------------------
# Privacy
# ---------------------------------------------------------------------------------------


def dump(engine: Any) -> str:
    with engine.connect() as connection:
        return "\n".join(connection.connection.iterdump())


def test_the_flow_stores_nothing_and_logs_no_candidate_data(
    flow: Flow, caplog: pytest.LogCaptureFixture, tmp_path: pathlib.Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    work, temp = tmp_path / "work", tmp_path / "temp"
    work.mkdir()
    temp.mkdir()
    monkeypatch.chdir(work)
    monkeypatch.setattr(tempfile, "tempdir", str(temp))
    before = dump(flow.engine)

    with caplog.at_level(logging.DEBUG):
        parsed = parse_resume(flow)
        profile = complete_profile(parsed)
        flow.client.post("/api/v1/candidates/validate", json=profile)
        flow.client.post("/api/v1/candidates/normalize", json=profile)
        body = recommend(flow, profile)
        export = flow.client.post(
            "/api/v1/applications/export", json={"records": tracker_records(flow, body)}
        )
    assert export.status_code == 200

    assert dump(flow.engine) == before, "no candidate data reached the database"
    assert set(Base.metadata.tables) == ALLOWED_OPERATIONAL_TABLES
    assert os.listdir(work) == [] and os.listdir(temp) == []

    logged = caplog.text + "\n".join(str(record.args) for record in caplog.records)
    for marker in MARKERS:
        assert marker not in logged, marker
        assert marker not in export.text
        assert marker not in json.dumps(body)


def test_invalid_payloads_in_the_flow_do_not_echo_candidate_data(flow: Flow) -> None:
    broken = {**PROFILE_B, "education": [{**PROFILE_B["education"][0], "cgpa": "not-a-number"}]}
    for url, payload in (
        ("/api/v1/candidates/validate", broken),
        ("/api/v1/recommendations", {"profile": broken}),
        ("/api/v1/applications/export", {"records": []}),
    ):
        response = flow.client.post(url, json=payload)
        assert response.status_code == 422
        assert response.json()["error"] == "VALIDATION_ERROR"
        assert "not-a-number" not in response.text
        assert "Mechanical Engineering" not in response.text
