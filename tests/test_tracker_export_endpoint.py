"""Tracker export API contract and privacy tests (PR 6A — ADR-023).

Drives ``POST /api/v1/applications/export`` through the real application. Proves the binary
contract, the unchanged error envelope, the OpenAPI description, and that the endpoint stores
nothing, reads nothing from the database, writes no file and logs only counts (INV-1, INV-4,
ADR-002).

All data is obviously synthetic.
"""

from __future__ import annotations

import io
import logging
import os
import pathlib
import re
import tempfile
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.models.job import Job, JobStatus
from app.schemas.candidate import JobType
from app.schemas.tracker import MAX_TRACKER_ROWS
from tests.conftest import ALLOWED_OPERATIONAL_TABLES

URL = "/api/v1/applications/export"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

JOB_MARKER = "JOB-MARKER-6a2e"
COMPANY = "Company Marker 6a2e Ltd"
ROLE = "Role Marker 6a2e Intern"
LINK = "https://link-marker-6a2e.example.com/apply"
REASON = "Reason marker 6a2e sentence."
NOTES = "Notes marker 6a2e: called the recruiter"
REQUIREMENT_TEXT = "Requirement marker 6a2e"
CANDIDATE_VALUE = "CandidateValue-6a2e"
REQUIREMENT_NOTE = "Requirement note marker 6a2e"
SCORE = 61.37
MARKERS = [JOB_MARKER, COMPANY, ROLE, LINK, REASON, NOTES, REQUIREMENT_TEXT, CANDIDATE_VALUE,
           REQUIREMENT_NOTE]


def requirement(**overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "requirement_type": "MIN_CGPA",
        "requirement": REQUIREMENT_TEXT,
        "candidate_value": CANDIDATE_VALUE,
        "status": "PASS",
        "confidence": "HIGH",
        "method": "deterministic",
        "reason_code": "MEETS_MINIMUM",
        "note": REQUIREMENT_NOTE,
    }
    base.update(overrides)
    return base


def marker_record(job_id: str = JOB_MARKER, **overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "job_id": job_id,
        "company_name": COMPANY,
        "role_title": ROLE,
        "apply_link": LINK,
        "deadline": "2026-11-15",
        "job_status": "ACTIVE",
        "eligibility_state": "ELIGIBLE",
        "match_score": SCORE,
        "reason": REASON,
        "application_status": "APPLIED",
        "evaluated_at": "2026-09-19T10:00:00Z",
        "notes": NOTES,
        "requirement_breakdown": [requirement()],
    }
    base.update(overrides)
    return base


def plain(job_id: str) -> dict[str, Any]:
    return {"job_id": job_id, "company_name": "Example Co", "role_title": "Intern"}


@dataclass
class Server:
    client: TestClient
    engine: Any
    queries: list[str]


@pytest.fixture
def server() -> Iterator[Server]:
    """The real application with an isolated catalogue, so any database use would be visible."""
    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    queries: list[str] = []

    @event.listens_for(engine, "before_cursor_execute")
    def _record(conn, cursor, statement, parameters, context, executemany):  # noqa: ANN001
        queries.append(statement)

    with factory() as seed:
        seed.add(Job(
            id="catalogue-job", company_name="Example Co", role_title="Intern",
            job_type=JobType.INTERNSHIP, description="Python.", requirements={},
            allowed_fields=[], required_skills=["Python"], source="test",
            content_hash="c" * 64, status=JobStatus.ACTIVE,
            last_verified_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        ))
        seed.commit()

    def override_get_db() -> Iterator[Session]:
        db = factory()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    queries.clear()
    try:
        yield Server(TestClient(app), engine, queries)
    finally:
        app.dependency_overrides.pop(get_db, None)
        engine.dispose()


def dump(server: Server) -> str:
    with server.engine.connect() as connection:
        return "\n".join(connection.connection.iterdump())


def open_workbook(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content))


# ---------------------------------------------------------------------------------------
# Contract
# ---------------------------------------------------------------------------------------


def test_export_returns_the_workbook(server: Server) -> None:
    response = server.client.post(URL, json={"records": [marker_record(), plain("second")]})
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == 'attachment; filename="eligicore-tracker.xlsx"'
    assert response.headers["cache-control"] == "no-store"
    assert response.content[:2] == b"PK"  # a zip container, not JSON
    book = open_workbook(response.content)
    assert book.sheetnames == ["Tracker", "Requirements"]
    tracker = book["Tracker"]
    assert [tracker.cell(row=r, column=2).value for r in (2, 3)] == [JOB_MARKER, "second"]
    assert tracker["G2"].value == SCORE and tracker["G3"].value is None
    assert tracker["M2"].value == NOTES
    assert book["Requirements"]["A2"].value == 1 and book["Requirements"]["B2"].value == JOB_MARKER


def test_request_id_is_preserved(server: Server) -> None:
    response = server.client.post(URL, json={"records": [plain("a")]},
                                  headers={"X-Request-ID": "req-export-1"})
    assert response.status_code == 200
    assert response.headers["x-request-id"] == "req-export-1"
    generated = server.client.post(URL, json={"records": [plain("a")]})
    assert generated.headers["x-request-id"]


def test_filename_never_carries_row_data(server: Server) -> None:
    response = server.client.post(URL, json={"records": [marker_record()]})
    disposition = response.headers["content-disposition"]
    for marker in MARKERS:
        assert marker not in disposition
    assert not re.search(r"\d{4}-?\d{2}-?\d{2}", disposition)  # no timestamp either


def test_five_hundred_rows_through_the_api(server: Server) -> None:
    records = [plain(f"job-{i:03d}") for i in range(MAX_TRACKER_ROWS)]
    response = server.client.post(URL, json={"records": records})
    assert response.status_code == 200
    assert open_workbook(response.content)["Tracker"].max_row == MAX_TRACKER_ROWS + 1


def test_formula_text_through_the_api_stays_text(server: Server) -> None:
    response = server.client.post(URL, json={"records": [plain("a") | {"notes": "=1+1"}]})
    cell = open_workbook(response.content)["Tracker"]["M2"]
    assert cell.value == "'=1+1" and cell.data_type == "s"


def test_get_is_not_allowed(server: Server) -> None:
    response = server.client.get(URL)
    assert response.status_code == 405
    assert response.json()["error"] == "HTTP_ERROR"


INVALID_PAYLOADS = [
    pytest.param({"records": []}, id="empty"),
    pytest.param({"records": [plain(f"j{i}") for i in range(MAX_TRACKER_ROWS + 1)]}, id="501-rows"),
    pytest.param({"records": [marker_record(), marker_record()]}, id="duplicate-job-id"),
    pytest.param({"records": [marker_record(application_status=f"STATUS {NOTES}")]}, id="bad-status"),
    pytest.param({"records": [marker_record(eligibility_state=COMPANY)]}, id="bad-eligibility"),
    pytest.param({"records": [marker_record(match_score=101)]}, id="score-above-100"),
    pytest.param({"records": [marker_record(match_score=-1)]}, id="score-negative"),
    pytest.param({"records": [marker_record(match_score=REASON)]}, id="score-text"),
    pytest.param({"records": [marker_record(notes=NOTES * 30)]}, id="oversized-notes"),
    pytest.param({"records": [marker_record(company_name=COMPANY * 10)]}, id="oversized-company"),
    pytest.param({"records": [marker_record(apply_link=LINK * 60)]}, id="oversized-link"),
    pytest.param({"records": [marker_record(requirement_breakdown=[
        requirement(note=REQUIREMENT_NOTE * 40)])]}, id="oversized-requirement"),
    pytest.param({"records": [marker_record(unknown=NOTES)]}, id="unknown-record-field"),
    pytest.param({"records": [marker_record()], "candidate_id": JOB_MARKER}, id="unknown-top-field"),
    pytest.param({"records": [marker_record(notes=NOTES + "\x00")]}, id="control-character"),
    pytest.param({"records": [marker_record(evaluated_at="2026-09-19T10:00:00")]}, id="naive-time"),
    pytest.param({"records": marker_record()}, id="wrong-type"),
    pytest.param({}, id="missing-records"),
]


@pytest.mark.parametrize("payload", INVALID_PAYLOADS)
def test_invalid_requests_are_422_in_the_standard_envelope_without_echo(
    server: Server, payload: dict[str, Any]
) -> None:
    response = server.client.post(URL, json=payload)
    assert response.status_code == 422
    assert response.headers["content-type"].startswith("application/json")
    body = response.json()
    assert set(body) == {"error", "message", "request_id", "details"}
    assert body["error"] == "VALIDATION_ERROR" and body["request_id"]
    assert body["request_id"] == response.headers["x-request-id"]
    for marker in MARKERS:
        assert marker not in response.text


def test_malformed_json_is_422(server: Server) -> None:
    response = server.client.post(
        URL, content=b'{"records": [', headers={"Content-Type": "application/json"}
    )
    assert response.status_code == 422
    assert response.json()["error"] == "VALIDATION_ERROR"


def test_unexpected_failure_is_a_generic_500_without_row_data(
    server: Server, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    def boom(records):  # noqa: ANN001, ANN202
        raise RuntimeError(f"failed on {records[0].notes} at {records[0].company_name}")

    monkeypatch.setattr("app.routers.applications.export_tracker", boom)
    client = TestClient(app, raise_server_exceptions=False)
    with caplog.at_level(logging.DEBUG):
        response = client.post(URL, json={"records": [marker_record()]})
    assert response.status_code == 500
    body = response.json()
    assert body["error"] == "INTERNAL_ERROR" and set(body) == {"error", "message", "request_id", "details"}
    for marker in MARKERS:
        assert marker not in response.text
        assert marker not in caplog.text


# ---------------------------------------------------------------------------------------
# OpenAPI and route surface
# ---------------------------------------------------------------------------------------


def test_openapi_documents_the_binary_response() -> None:
    spec = TestClient(app).get("/openapi.json").json()
    operations = spec["paths"][URL]
    assert set(operations) == {"post"}
    post = operations["post"]
    assert post["tags"] == ["applications"]
    assert post["requestBody"]["content"]["application/json"]["schema"]["$ref"].endswith(
        "/TrackerExportRequest"
    )
    ok = post["responses"]["200"]["content"]
    assert list(ok) == [XLSX]
    assert ok[XLSX]["schema"] == {"type": "string", "format": "binary"}
    assert "422" in post["responses"]
    status_enum = spec["components"]["schemas"]["ApplicationStatus"]["enum"]
    assert status_enum == ["NOT_APPLIED", "APPLIED", "INTERVIEW", "REJECTED", "OFFER"]


def test_deferred_endpoints_do_not_exist(server: Server) -> None:
    spec = server.client.get("/openapi.json").json()
    for path in ("/api/v1/matching/score", "/api/v1/jobs/ingest"):
        assert path not in spec["paths"], path
    assert server.client.post("/api/v1/matching/score", json={}).status_code == 404
    # /jobs/ingest has no route of its own: the path only matches GET /jobs/{job_id}.
    assert server.client.post("/api/v1/jobs/ingest", json={}).status_code == 405

    # Application preparation stopped being deferred in Slice 7B (ADR-025) and is the only
    # endpoint added with it. It is a POST, and it is not a second binary response.
    assert set(spec["paths"]["/api/v1/applications/prepare"]) == {"post"}
    assert "application/json" in (
        spec["paths"]["/api/v1/applications/prepare"]["post"]["responses"]["200"]["content"]
    )


def test_openapi_structure_is_unchanged_by_the_status_text() -> None:
    """Week 9A-1 corrected prose only, so the contract surface is pinned exactly here.

    Paths, methods, response status codes and component schema names are the whole of what a
    client depends on. Pinning them means a future edit to the description — the one part of
    the document Week 9 is allowed to touch — cannot quietly move the contract with it
    (ADR-027 D10).
    """
    spec = TestClient(app).get("/openapi.json").json()

    paths = (
        "/api/v1/applications/export", "/api/v1/applications/prepare",
        "/api/v1/candidates/normalize", "/api/v1/candidates/validate",
        "/api/v1/eligibility/check", "/api/v1/health", "/api/v1/jobs", "/api/v1/jobs/{job_id}",
        "/api/v1/recommendations", "/api/v1/resumes/parse",
    )
    assert set(spec["paths"]) == set(paths)
    assert len(spec["paths"]) == 10

    operations = {
        (path, method): sorted(operation["responses"])
        for path, methods in spec["paths"].items()
        for method, operation in methods.items()
    }
    assert operations == {
        ("/api/v1/applications/export", "post"): ["200", "422"],
        ("/api/v1/applications/prepare", "post"): ["200", "404", "422"],
        ("/api/v1/candidates/normalize", "post"): ["200", "422"],
        ("/api/v1/candidates/validate", "post"): ["200", "422"],
        ("/api/v1/eligibility/check", "post"): ["200", "422"],
        ("/api/v1/health", "get"): ["200"],
        ("/api/v1/jobs", "get"): ["200", "422"],
        ("/api/v1/jobs/{job_id}", "get"): ["200", "404", "422"],
        ("/api/v1/recommendations", "post"): ["200", "422"],
        ("/api/v1/resumes/parse", "post"): ["200", "413", "415", "422", "502"],
    }

    schemas = (
        "AnswerOutcome", "ApplicationPrepareRequest", "ApplicationPrepareResponse",
        "ApplicationQuestion", "ApplicationStatus",
        "Body_parse_resume_api_v1_resumes_parse_post", "CandidateNormalizationResponse",
        "CandidatePreferences", "CandidateProfile-Input", "CandidateProfile-Output",
        "CandidateValidationResponse", "CertificationEntry", "ClaimCategory", "Confidence",
        "DegreeLevel", "EducationEntry", "EligibilityCheckRequest", "EligibilityCheckResponse",
        "EligibilityState", "EvaluationMethod", "ExperienceEntry", "ExtractionMetadata",
        "GenerationOutcome", "GradeScale", "HTTPValidationError", "HealthResponse",
        "IssueSeverity", "JobEligibility", "JobListResponse", "JobRead", "JobStatusSchema",
        "JobType", "MatchScoreBasis", "MatchTermKind", "NormalizedGrade", "PackageStatus",
        "PreparedAnswer", "ProjectEntry", "ReasonCode", "RecommendationItem",
        "RecommendationMatch", "RecommendationRequest", "RecommendationResponse",
        "RemovalReason", "RemovalScope", "RemovedClaim", "RequirementProvenance",
        "RequirementResult", "RequirementStatus", "RequirementType", "ResumeParseResponse",
        "ResumeParseStatus",
        "SharedTerm", "SourceFormat", "TrackerExportRequest", "TrackerRecord",
        "ValidationError", "ValidationIssue", "WorkMode",
    )
    assert set(spec["components"]["schemas"]) == set(schemas)
    assert len(spec["components"]["schemas"]) == 59


def test_openapi_status_text_is_current() -> None:
    description = TestClient(app).get("/openapi.json").json()["info"]["description"]
    assert "Weeks 1–8 of a 10-week build are complete." in description
    assert "tracker export" in description
    for stale in (
        "Week 5 of a 10-week build is in progress",
        "export are not yet implemented",
        "Weeks 1–5",
        "Weeks 1–6",
        "Weeks 1–7",
        "Week 6 is underway",
    ):
        assert stale not in description, stale
    for premature in ("Weeks 1–9", "Week 9", "Week 10", "10-week build is complete"):
        assert premature not in description, premature


# ---------------------------------------------------------------------------------------
# Privacy — nothing stored, nothing written, nothing logged
# ---------------------------------------------------------------------------------------


def test_export_makes_no_database_query_and_changes_nothing(server: Server) -> None:
    before = dump(server)
    server.queries.clear()
    assert server.client.post(URL, json={"records": [marker_record()]}).status_code == 200
    assert server.client.post(URL, json={"records": []}).status_code == 422
    assert server.queries == []
    assert dump(server) == before
    assert set(Base.metadata.tables) == ALLOWED_OPERATIONAL_TABLES


def test_export_writes_no_file(
    server: Server, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work, temp = tmp_path / "work", tmp_path / "temp"
    work.mkdir()
    temp.mkdir()
    monkeypatch.chdir(work)
    monkeypatch.setattr(tempfile, "tempdir", str(temp))
    repo_files = sorted(p.name for p in pathlib.Path(__file__).resolve().parents[1].iterdir())
    response = server.client.post(URL, json={"records": [marker_record(), plain("b")]})
    assert response.status_code == 200
    assert os.listdir(work) == [] and os.listdir(temp) == []
    assert sorted(p.name for p in pathlib.Path(__file__).resolve().parents[1].iterdir()) == repo_files


def test_only_a_counts_line_is_logged(server: Server, caplog: pytest.LogCaptureFixture) -> None:
    with caplog.at_level(logging.DEBUG):
        response = server.client.post(URL, json={"records": [marker_record(), plain("b")]})
    assert response.status_code == 200
    logged = caplog.text + "\n".join(str(record.args) for record in caplog.records)
    for marker in MARKERS + [str(SCORE), "2026-11-15"]:
        assert marker not in logged, marker
    lines = [r.getMessage() for r in caplog.records if r.name == "eligicore.tracker_export"]
    assert len(lines) == 1
    assert re.fullmatch(
        r"tracker_export rows=2 requirement_rows=1 bytes=\d+ duration_ms=[\d.]+", lines[0]
    )
    assert f"bytes={len(response.content)} " in lines[0]


def test_invalid_request_logs_nothing_from_the_body(
    server: Server, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level(logging.DEBUG):
        server.client.post(URL, json={"records": [marker_record(), marker_record()]})
        server.client.post(URL, json={"records": [marker_record(application_status=NOTES)]})
    for marker in MARKERS:
        assert marker not in caplog.text
